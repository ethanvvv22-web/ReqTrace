"""Order-placement documentation.

Produces the requirement specification for a case and, once supplier answers
are assessed, a compliance matrix and a deviation list. Every row carries the
requirement id + version, the source reference, the supplier's own words, how
the status was reached (rule / ai) and the reviewer's decision.
"""
from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from .model import Assessment, ProcurementCase, SupplierAnswer
from .select import FeatureHit, Selected

MATRIX_COLS = ["req_id", "version", "title", "source_ref", "supplier_claimed",
               "supplier_answer", "status", "evidence", "assessed_by", "confidence",
               "flags", "reviewer_decision", "reviewer_comment"]


def spec_markdown(case: ProcurementCase, hits: list[FeatureHit],
                  selected: list[Selected]) -> str:
    lines = [f"# Requirement specification - {case.title} ({case.id})",
             f"_Generated {date.today().isoformat()} by ReqTrace. Draft for engineer review._", "",
             "## Case", case.description, "", "## Machine features identified", "",
             "| Feature | Found by | Evidence |", "|---|---|---|"]
    lines += [f"| {h.feature} | {h.source} | \"{h.evidence}\" |" for h in hits]
    lines += ["", f"## Requirements ({len(selected)})", "",
              "| ID | v | Requirement | Source | Applies because | Verification |",
              "|---|---|---|---|---|---|"]
    for s in selected:
        r = s.req
        why = ", ".join(s.because) or "always"
        lines.append(f"| {r.id} | {r.version} | **{r.title}.** {r.text} | {r.source_ref} "
                     f"| {why} | {r.verification or '-'} |")
    return "\n".join(lines) + "\n"


def matrix_rows(answers: list[SupplierAnswer], assessments: list[Assessment],
                library: dict, decisions: dict | None = None) -> list[dict]:
    decisions = decisions or {}
    rows = []
    for ans, a in zip(answers, assessments):
        r = library.get(ans.req_id)
        d = decisions.get(ans.req_id, {})
        rows.append({
            "req_id": ans.req_id, "version": r.version if r else "",
            "title": r.title if r else "UNKNOWN REQUIREMENT", "source_ref": r.source_ref if r else "",
            "supplier_claimed": ans.claimed, "supplier_answer": ans.answer,
            "status": d.get("status", a.status.value), "evidence": a.evidence,
            "assessed_by": "human" if d else a.layer, "confidence": a.confidence,
            "flags": "; ".join(a.flags), "reviewer_decision": d.get("action", ""),
            "reviewer_comment": d.get("comment", ""),
        })
    return rows


def write_matrix(rows: list[dict], out: Path) -> list[Path]:
    """Writes CSV always, and XLSX too if openpyxl is installed."""
    out.parent.mkdir(parents=True, exist_ok=True)
    written = []
    csv_path = out.with_suffix(".csv")
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=MATRIX_COLS)
        w.writeheader()
        w.writerows(rows)
    written.append(csv_path)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return written
    wb = Workbook()
    ws = wb.active
    ws.title = "Compliance matrix"
    ws.append(MATRIX_COLS)
    for c in ws[1]:
        c.font = Font(bold=True)
    fills = {"non_compliant": "F4CCCC", "partial": "FCE5CD", "unclear": "FFF2CC"}
    for row in rows:
        ws.append([row[k] for k in MATRIX_COLS])
        color = fills.get(row["status"])
        if color:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=color)
    dev = wb.create_sheet("Deviations")
    dev.append(MATRIX_COLS)
    for row in rows:
        if row["status"] != "compliant":
            dev.append([row[k] for k in MATRIX_COLS])
    xlsx_path = out.with_suffix(".xlsx")
    wb.save(xlsx_path)
    written.append(xlsx_path)
    return written


def deviations_markdown(rows: list[dict]) -> str:
    dev = [r for r in rows if r["status"] != "compliant"]
    lines = [f"# Deviation list ({len(dev)} of {len(rows)} requirements)", "",
             "| ID | Requirement | Status | Supplier said | Flags |", "|---|---|---|---|---|"]
    lines += [f"| {r['req_id']} | {r['title']} | {r['status']} | {r['supplier_answer']} "
              f"| {r['flags'] or '-'} |" for r in dev]
    return "\n".join(lines) + "\n"
